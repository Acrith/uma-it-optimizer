"""Write uma-it-web's loop sheet data from the community's looping spreadsheet:
per running style (Front, Pace, Late, End), each skill with its tier and
whether it is out in Global yet. The site builds its built-in target sets
from it (collection page).

The sheet: one tab per style; column A holds a skill or a tier header
(grey fill, or text like "T1"); a pink fill (F4CCCC) marks a skill not
released in Global. The first line of column A is the skill's name as the
community writes it (the site maps it to the game's spark names).

    uv run --no-project --with openpyxl python export_loop_sheet.py [xlsx]

Without a file it downloads the sheet's xlsx export.
"""
from __future__ import annotations

import json
import re
import sys
import urllib.request
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

SHEET_ID = "1pWBnsoH2h7NDm9Q5ZRZ-s9JuyMN-aAI5wJbIEqACCt4"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"
EXPORT = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
OUT = Path(__file__).parent / "../../../uma-it-web/uma_it_web/enrich/data/loop_sheet.json"
STYLES = ("Front", "Pace", "Late", "End")
HEADER_FILL = "FFD9D9D9"
UNRELEASED_FILL = "FFF4CCCC"
# A tier header even where the grey fill is missing (Late's "T1").
TIER_TEXT = re.compile(r"^T\d+(\.\d+)?$")


def fill(cell) -> str | None:
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


def read(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    styles = {}
    for style in STYLES:
        ws = wb[style]
        tier, rows = None, []
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, 1)
            if cell.value is None:
                continue
            text = str(cell.value).split("\n")[0].strip()
            if fill(cell) == HEADER_FILL or TIER_TEXT.match(text):
                tier = text
                continue
            rows.append({"tier": tier, "skill": text, "unreleased": fill(cell) == UNRELEASED_FILL})
        styles[style] = rows
    return styles


def main() -> int:
    if len(sys.argv) > 1:
        path = Path(sys.argv[1])
    else:
        path = Path("/tmp/loop_sheet.xlsx")
        urllib.request.urlretrieve(EXPORT, path)
    styles = read(path)
    fetched = datetime.now(UTC).isoformat(timespec="seconds")
    out = {"_meta": {"source": SHEET_URL, "fetched_at": fetched}, "styles": styles}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    for style, rows in styles.items():
        tiers = sorted({r["tier"] for r in rows if r["tier"]})
        print(f"{style}: {len(rows)} rows, tiers {', '.join(tiers)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
